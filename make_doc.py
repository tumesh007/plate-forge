"""Build the Plate Forge project document."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
DOCS = ROOT / "docs"
OUT = ROOT / "Plate-Forge-Documentation.docx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x63, 0x6E)
ACCENT = RGBColor(0x0E, 0x7C, 0x5C)

doc = Document()

# ---- page + base styles -------------------------------------------------
sec = doc.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.4)
sec.right_margin = Cm(2.4)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.18
# Arial must also be set for East Asian fallback
rpr = normal.element.get_or_add_rPr()
rfonts = rpr.find(qn("w:rFonts"))
if rfonts is None:
    rfonts = OxmlElement("w:rFonts")
    rpr.append(rfonts)
rfonts.set(qn("w:eastAsia"), "Arial")

for name, size, color, space_before in [
    ("Heading 1", 15, INK, 16),
    ("Heading 2", 11.5, ACCENT, 12),
]:
    st = doc.styles[name]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(space_before)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.keep_with_next = True


def p(text="", style=None, size=None, color=None, bold=False, italic=False, align=None):
    para = doc.add_paragraph(style=style)
    run = para.add_run(text)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    if align is not None:
        para.alignment = align
    return para


def bullet(text, bold_lead=None):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = para.add_run(bold_lead)
        r.bold = True
    para.add_run(text)
    return para


def rule():
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(10)
    pPr = para._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C9D1D9")
    borders.append(bottom)
    pPr.append(borders)


def picture(name, width_cm, caption):
    path = DOCS / name
    if not path.exists():
        print("missing image:", path)
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(3)
    para.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(caption)
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    r.italic = True


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.name = "Arial"
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:fill"), "EEF3F1")
        hdr[i]._tc.get_or_add_tcPr().append(shade)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
            run.font.name = "Arial"
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ---- title block --------------------------------------------------------
title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(2)
run = title.add_run("Plate Forge")
run.font.size = Pt(26)
run.font.bold = True
run.font.color.rgb = INK
run.font.name = "Arial"

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(2)
run = sub.add_run("A single-sided car parking number plate you can re-letter in fifteen minutes")
run.font.size = Pt(12.5)
run.font.color.rgb = MUTED

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(14)
run = meta.add_run("Design note  ·  print settings  ·  OpenSCAD source  ·  October 2026")
run.font.size = Pt(9)
run.font.color.rgb = MUTED

rule()

picture("assembled.png", 15.4, "The panel seated in the easel at a 30° rake. The number is a 0.6 mm engraved groove.")

# ---- 1 ------------------------------------------------------------------
doc.add_heading("1. What this is", level=1)
p(
    "A parking plate consists of two parts. The number panel is a flat 3 mm plate with the digits "
    "engraved 0.6 mm into the face; it prints face up, flat on the bed, with no supports. The easel "
    "is a separate wedge that holds the panel at a 30° rake and is printed once. Because the easel is "
    "reusable, changing the number costs one panel — about fifteen minutes and eleven grams — instead "
    "of a whole new plate."
)
p(
    "Everything is sized for a Bambu Lab A1 mini. With the default eleven-character number the panel "
    "comes out 158 × 30 × 3 mm and the easel 167 × 28 × 32 mm, so both fit the 180 mm bed together."
)

# ---- 2 ------------------------------------------------------------------
doc.add_heading("2. The decision that shapes everything", level=1)
p(
    "Two reasonable requirements pull against each other, and the choice between them decides the "
    "rest of the design: PLA is the plastic that takes ink, and PETG is the plastic that survives a "
    "car. You cannot have both on this machine."
)
table(
    ["", "PLA", "PETG"],
    [
        ["Marker / paint adhesion", "good; forgiving of a missed primer", "poor; ink beads on it, primer is not optional"],
        ["Softens at", "57–60 °C", "80–85 °C"],
        ["Dashboard in an Indian summer", "70–80 °C — sags over a few hot afternoons", "marginal, but it survives"],
        ["Use it when", "the plate lives indoors or in shade", "the plate lives in the car"],
    ],
    widths=[4.4, 5.3, 5.7],
)
p(
    "ASA would settle the argument, but an A1 mini cannot print it — it is an open-frame machine with "
    "an 80 °C bed, and ASA wants 90–110 °C. So the default is PETG, and the glue-wash step in section 5 "
    "is what buys back the ink adhesion."
)

doc.add_heading("Engraved, not raised", level=2)
p(
    "A groove has two properties a raised ridge does not. Its rim acts as a paint dam, so you can flood "
    "the groove and wipe the face clean without the ink wandering. And its floor is a horizontal top "
    "surface, which means the slicer can iron it flat and seal the layer lines — the paths ink would "
    "otherwise creep along. Vertical groove walls also mean no layer stair-stepping, so the digit "
    "outline is exactly as sharp as the model."
)
p(
    "Raised digits are still available in both the generator and the OpenSCAD file. They suit a "
    "chisel-tip pen laid almost flat and wiped across the tops, but the edge is less forgiving."
)

# ---- 3 ------------------------------------------------------------------
doc.add_heading("3. What you get", level=1)
table(
    ["Part", "Size (mm)", "Volume", "Print once?", "Notes"],
    [
        ["Number panel", "158 × 30 × 3", "13.6 cm³", "per number", "engraved digits, prints face up"],
        ["Easel", "167 × 28 × 32", "17.4 cm³", "yes, reusable", "prints standing, 20° overhang, no supports"],
        ["Pen test coupon", "164 × 23 × 3.3", "7.9 cm³", "once", "six combinations in one short print"],
    ],
    widths=[3.2, 3.2, 2.2, 2.4, 4.4],
)
p(
    "Volumes are the solid model volume. Actual filament use is lower — a 3 mm plate is nearly all walls "
    "and skins, so budget roughly 11 g of panel and 7 g of easel at 15 % infill."
)

picture("detail.png", 14.6, "The groove at 0.6 mm deep, filled with a paint pen. The rim does the masking.")

# ---- 4 ------------------------------------------------------------------
doc.add_heading("4. Printing", level=2)
p("Bambu Studio names; OrcaSlicer uses the same ones.", italic=True, size=9.5, color=MUTED)
table(
    ["Setting", "Value", "Why"],
    [
        ["Layer height", "0.20 mm", "0.6 mm groove = exactly three layers"],
        ["Walls / top / bottom", "4 / 5 / 4", "the groove floor must be a solid skin"],
        ["Infill", "15 % gyroid", ""],
        ["Elephant-foot compensation", "0.15 mm", "otherwise the panel edge fights the easel slot"],
        ["Ironing type", "All top surfaces", "see the warning below"],
        ["Ironing flow / spacing", "10 % / 0.10 mm", "closes the inter-layer grooves"],
        ["Ironing speed / inset", "30 mm/s / 0.2 mm", ""],
        ["Plate", "textured PEI + glue (PETG)", "PETG bonds to smooth PEI hard enough to tear it"],
    ],
    widths=[4.6, 4.0, 6.8],
)
warning = doc.add_paragraph()
warning.paragraph_format.space_before = Pt(4)
r = warning.add_run("The one setting people get wrong: ")
r.bold = True
warning.add_run(
    "ironing must be set to \u201cAll top surfaces\u201d, not \u201cOnly topmost\u201d. The groove floors sit "
    "0.6 mm below the plate rim, so \u201ctopmost\u201d irons the rim, skips all ten digits, and wastes the print."
)

picture("print-plate.png", 15.4, "Both parts on the plate, as exported. The panel prints flat, the easel standing.")

# ---- 5 ------------------------------------------------------------------
doc.add_heading("5. Painting the number", level=1)
p(
    "Ink wicking along layer lines is a real, documented effect and it is capillary, not sloppiness. "
    "The cheap fix that works is a diluted PVA wash in the groove before you flood it."
)
for step in [
    ("Wipe the face with 70 % IPA. ", "Skin oil is the usual reason ink \u201cdoesn't stick\u201d."),
    ("Brush a thin PVA or wood-glue wash into the grooves. ", "Let it dry completely. This is the step that stops the creep."),
    ("Flood with an oil-based paint pen. ", "Sharpie Oil-Based or Posca PC-5M. Overfill slightly; do not try to stay inside the lines."),
    ("Wipe the flat face while the ink is tacky. ", "A barely damp cloth over a fingertip. The rim does the masking."),
    ("Cure 24 hours, then two thin matte UV clear coats. ", "Matte, because gloss makes ironing marks visible."),
]:
    bullet(step[1], bold_lead=step[0])

rule()
picture("coupon.png", 15.0, "The pen test coupon: both relief styles at 0.4, 0.6 and 0.9 mm. Print this first.")

# ---- 6 ------------------------------------------------------------------
doc.add_heading("6. The OpenSCAD file", level=1)
p(
    "openscad/plate_forge.scad is written for MakerWorld's Parametric Model Maker, not just for the "
    "desktop. It uses Customizer-style parameters with dropdowns at the top of the file, exposes the "
    "printable output through mw_plate_1() and an assembled preview through mw_assembly_view(), and "
    "pulls in no external libraries."
)
table(
    ["Parameter", "Default", "What it does"],
    [
        ["number", "98765 43210", "up to 14 characters; digits, space, + - ."],
        ["mode", "both", "panel + easel, panel only, easel only, or coupon"],
        ["lettering", "engraved", "engraved or raised"],
        ["relief_depth", "0.6 mm", "groove depth / ridge height"],
        ["digit_height", "18 mm", "cap height; shrinks automatically to fit the bed"],
        ["stroke_ratio", "0.13", "stroke as a fraction of digit height"],
        ["panel_thickness", "3.0 mm", "3 mm is the minimum for PETG on a dashboard"],
        ["rake_deg", "30°", "easel tilt from vertical"],
    ],
    widths=[4.2, 3.4, 7.8],
)
p(
    "The layout is a 7-segment skeleton with real gaps between every bar, because two bars that touch, "
    "share an edge, or meet at a single corner make a self-intersecting sketch, and OpenSCAD's "
    "difference() will quietly drop geometry rather than complain. Every segment is separated by "
    "exactly segment_gap."
)

# ---- 7 ------------------------------------------------------------------
doc.add_heading("7. The web generator", level=1)
p(
    "The same geometry exists as a browser tool at tumesh007.github.io/plate-forge. It takes a number, "
    "renders the plate live, and exports STLs — panel, easel, both on one plate, or the coupon. There is "
    "no build step and no CDN dependency; it runs from a static folder."
)
p(
    "The browser version builds the panel as a single closed manifold surface rather than two solids "
    "stacked, and its test suite asserts that every triangle edge is shared by exactly two triangles "
    "across twelve configurations. That check exists because an interior face sitting exactly on the "
    "groove floor can be read by a slicer as a filled cavity, which swallows the digits without any "
    "error message."
)

picture("easel.png", 14.2, "The easel as printed: two side fins, a plinth that doubles as the front lip, and a rear tie.")

# ---- 8 ------------------------------------------------------------------
doc.add_heading("8. Files", level=1)
table(
    ["Path", "What it is"],
    [
        ["openscad/plate_forge.scad", "the parametric model, MakerWorld ready"],
        ["index.html, js/, css/", "the browser generator"],
        ["OPTIONS.md", "the research: what the alternatives are and why this one won"],
        ["PRINT-SETTINGS.md", "slicer settings and the ink workflow in full"],
        ["MAKERWORLD.md", "listing copy, tags and settings for the upload"],
        ["test/check.mjs", "headless geometry tests"],
        ["docs/", "the images in this document"],
    ],
    widths=[6.2, 9.2],
)

p(
    "Print the coupon before the plate. It costs eight minutes and it turns every claim above into a "
    "measurement on your own filament.",
    italic=True,
    color=MUTED,
)

doc.save(OUT)
print("saved", OUT)

